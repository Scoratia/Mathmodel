"""Read-only independent checks of the exported Excel files against solver payloads."""
import json,math,hashlib,sys
import openpyxl
import numpy as np
from model import ROOT


def main():
    res=ROOT/'results'
    v=json.loads((res/'validation.json').read_text(encoding='utf-8'))
    checks=[]
    inp=json.loads((ROOT/'data/inputs.json').read_text(encoding='utf-8'))
    for name in ['ambient','radius']:
        a=np.asarray(inp[name],float)
        assert np.isfinite(a).all() and a[0,0]==0 and np.all(np.diff(a[:,0])>0)
    assert np.all(np.diff(np.asarray(inp['radius'])[:,1])<=0)
    for q in ['1','3','4']:
        b=v['balance'][q]
        assert b['C_min']>=0 and b['C_max']<=2.55+1e-8
        assert b['max_radial_C_increase']<1e-8
        assert b['relative_to_initial_C']<2e-5
    assert v['analytic'][-1]['equivalent_22K_step_error_C']<3e-5
    assert abs(v['time_solver']['difference_s'])<.1
    for q in [3,4]:
        row=[r for r in v['mesh'] if r['problem']==q][-1]
        assert abs(row['delta_time_s'])<.2
        assert v['crossing'][str(q)]['max_C_first_full_minute']<.15
    for q in range(1,5):
        f=ROOT/f'result{q}.xlsx'
        payload=json.loads((res/f'result{q}.json').read_text(encoding='utf-8'))
        w=openpyxl.load_workbook(f,data_only=True,read_only=True)
        assert w.sheetnames==list(payload)
        count=0;blanks=0
        for name,expected in payload.items():
            s=w[name]
            if s.max_row is None or s.max_column is None:s.calculate_dimension(force=True)
            assert (s.max_row,s.max_column)==(len(expected),len(expected[0]))
            for i,(row,exp) in enumerate(zip(s.values,expected)):
                assert len(row)==len(exp)
                for j,(actual,want) in enumerate(zip(row,exp)):
                    if want is None:
                        assert actual is None,(q,name,i,j,actual,want);blanks+=1
                    elif isinstance(want,(float,int)):
                        assert isinstance(actual,(float,int)) and math.isfinite(actual)
                        assert abs(actual-want)<1e-9,(q,name,i,j,actual,want)
                    else:assert actual==want,(q,name,i,j,actual,want)
                    count+=1
                if i>0:
                    assert row[0]==i*(1 if q<=2 else 60)
                    if q==4:
                        R=row[-1]
                        for j,r in enumerate(np.arange(0,2.001,.1),start=1):
                            if r>R+1e-6:assert row[j] is None
                            elif r<R-1e-6:assert row[j] is not None
            assert s.cell(2,2).number_format=='0.0000'
        checks.append({'file':f.name,'checked_cells':count,'blank_outside_material':blanks,
                       'sha256':hashlib.sha256(f.read_bytes()).hexdigest(),'status':'PASS'})
        w.close()
    arrays=np.load(res/'profiles_full_precision.npz')
    assert all(np.isfinite(arrays[k]).all() for k in arrays.files)
    out={'numerical_gate':'PASS','export_gate':'PASS','workbooks':checks,
         'scope':'All saved workbook cells compared to rounded solver payloads; numerical tests concern stated effective model, not experimental validation.'}
    (res/'delivery_checks.json').write_text(json.dumps(out,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(out,ensure_ascii=False,indent=2))


if __name__=='__main__':main()
