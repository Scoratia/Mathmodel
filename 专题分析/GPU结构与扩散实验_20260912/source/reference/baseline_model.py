"""径向药材烘干模型（SI 单位，温度状态为摄氏度）。

本模块复现题给有效热湿扩散模型。第 4 问采用长度不变、均匀仿射
径向收缩；C 是干基质量比，因此没有额外“浓缩项”。可选表面潜热
情景统一采用干固体守恒定义的出水通量，但仍保留题给有效热容量。
"""
from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import Any
import hashlib
import math

import numpy as np
from scipy.integrate import solve_ivp
from scipy.interpolate import PchipInterpolator
from scipy.sparse import bmat, diags
from openpyxl import load_workbook


R0 = 0.02
LENGTH = 0.25
T0 = 28.0
C0 = 2.55
H = 25.0
HM = 8.0e-7
THRESHOLD = 0.15


@dataclass
class ModelInputs:
    """原始附件数值及边界插值；所有时间 s，半径 m。"""

    boundary_time: np.ndarray
    ambient_temperature: np.ndarray
    ambient_moisture: np.ndarray
    radius_time: np.ndarray
    radius_values: np.ndarray
    source_hashes: dict[str, str] = field(default_factory=dict)

    def __post_init__(self):
        names = ("boundary_time", "ambient_temperature", "ambient_moisture",
                 "radius_time", "radius_values")
        for name in names:
            values = np.asarray(getattr(self, name), dtype=float)
            if values.ndim != 1 or not np.all(np.isfinite(values)):
                raise ValueError(f"{name} 必须是一维有限数值数组")
            setattr(self, name, values.copy())
        if not (len(self.boundary_time) == len(self.ambient_temperature)
                == len(self.ambient_moisture)):
            raise ValueError("烘房时间、温度和浓度长度不一致")
        if len(self.radius_time) != len(self.radius_values):
            raise ValueError("半径时间与数值长度不一致")
        for name in ("boundary_time", "radius_time"):
            values = getattr(self, name)
            if len(values) < 2 or values[0] != 0 or np.any(np.diff(values) <= 0):
                raise ValueError(f"{name} 须从 0 开始且严格递增")
        if np.any(self.radius_values <= 0):
            raise ValueError("半径必须为正")
        if np.any(np.diff(self.radius_values) > 1e-12):
            raise ValueError("实测半径不是单调不增；请检查原数据，程序不会静默修改")
        if not np.isclose(self.radius_values[0], R0, rtol=0, atol=1e-10):
            raise ValueError("附件初始半径与题给 0.02 m 不一致")
        if np.any(self.ambient_moisture < 0):
            raise ValueError("环境含水率不能为负")
        self._radius_interpolator = PchipInterpolator(
            self.radius_time, self.radius_values, extrapolate=False)

    @classmethod
    def from_excel(cls, directory: str | Path) -> "ModelInputs":
        """读取附件1.xlsx 和附件2.xlsx，半径列由 cm 转为 m。

        directory 可指向附件文件夹、A题文件夹，或该题所在工作区。
        不搜索用户目录，也不隐式使用其他版本数据。
        """
        base = Path(directory).expanduser().resolve()
        if base.is_file():
            base = base.parent
        candidates = (base, base / "附件", base / "A题" / "附件")
        actual = next((p for p in candidates
                       if (p / "附件1.xlsx").is_file()
                       and (p / "附件2.xlsx").is_file()), None)
        if actual is None:
            raise FileNotFoundError(f"未在 {base} 下找到附件1.xlsx、附件2.xlsx")

        def read_numbers(path: Path, width: int) -> np.ndarray:
            workbook = load_workbook(path, read_only=True, data_only=True)
            try:
                rows = list(workbook.active.iter_rows(values_only=True))
                data = []
                for rownum, row in enumerate(rows[1:], start=2):
                    values = row[:width]
                    if all(v is None for v in values):
                        continue
                    if len(values) != width or any(v is None for v in values):
                        raise ValueError(f"{path.name} 第 {rownum} 行存在空值")
                    try:
                        data.append([float(v) for v in values])
                    except (TypeError, ValueError) as exc:
                        raise ValueError(f"{path.name} 第 {rownum} 行不是数值") from exc
                if not data:
                    raise ValueError(f"{path.name} 没有数据行")
                return np.asarray(data, dtype=float)
            finally:
                workbook.close()

        boundary_path, radius_path = actual / "附件1.xlsx", actual / "附件2.xlsx"
        boundary_data = read_numbers(boundary_path, 3)
        radius_data = read_numbers(radius_path, 2)
        hashes = {str(p): hashlib.sha256(p.read_bytes()).hexdigest()
                  for p in (boundary_path, radius_path)}
        return cls(boundary_data[:, 0], boundary_data[:, 1], boundary_data[:, 2],
                   radius_data[:, 0], radius_data[:, 1] / 100.0, hashes)

    @property
    def boundary_end(self) -> float:
        return float(self.boundary_time[-1])

    @property
    def radius_end(self) -> float:
        return float(self.radius_time[-1])

    @property
    def plateau(self) -> tuple[float, float]:
        """最后一小时采样点的算术均值（本题包括两端共 61 点）。"""
        mask = self.boundary_time >= self.boundary_end - 3600.0
        return (float(np.mean(self.ambient_temperature[mask])),
                float(np.mean(self.ambient_moisture[mask])))

    def boundary(self, t, extension="last_hour_mean", phase=None):
        """前 4 h 线性插值，之后采用已声明的平台外推假设。

        phase 仅供分段求解：observed/plateau 确定断点所在段的左右值。
        默认断点本身取最后一个实测值。没有未说明的平滑处理。
        """
        t = np.asarray(t, dtype=float)
        if np.any(~np.isfinite(t)) or np.any(t < 0):
            raise ValueError("时间必须为非负有限数")
        if extension == "last_hour_mean":
            ta_end, ca_end = self.plateau
        elif extension == "last_value":
            ta_end, ca_end = self.ambient_temperature[-1], self.ambient_moisture[-1]
        else:
            raise ValueError("extension 须为 last_hour_mean 或 last_value")
        if phase not in (None, "observed", "plateau"):
            raise ValueError("未知边界分段")
        ta = np.interp(t, self.boundary_time, self.ambient_temperature)
        ca = np.interp(t, self.boundary_time, self.ambient_moisture)
        use_plateau = t > self.boundary_end
        if phase == "plateau":
            use_plateau = np.ones_like(t, dtype=bool)
        elif phase == "observed":
            use_plateau = np.zeros_like(t, dtype=bool)
        ta, ca = np.where(use_plateau, ta_end, ta), np.where(use_plateau, ca_end, ca)
        if t.ndim == 0:
            return float(ta), float(ca)
        return ta, ca

    def radius(self, t):
        t = np.asarray(t, dtype=float)
        if np.any(~np.isfinite(t)) or np.any(t < 0) or np.any(t > self.radius_end):
            raise ValueError(f"半径只允许在实测区间 [0, {self.radius_end:g}] s 内插值")
        r = self._radius_interpolator(t)
        return float(r) if t.ndim == 0 else r


@dataclass
class FVMGrid:
    """N 个径向区间，N+1 个节点（中心和表面均显式保留）。"""

    N: int
    grid_kind: str = "uniform"
    cluster_power: float = 1.5

    def __post_init__(self):
        if not isinstance(self.N, (int, np.integer)) or self.N < 4:
            raise ValueError("N 必须为不小于 4 的整数")
        s = np.linspace(0.0, 1.0, self.N + 1)
        if self.grid_kind == "uniform":
            self.x = s
        elif self.grid_kind == "surface_clustered":
            if not 1.0 <= self.cluster_power <= 2.0:
                raise ValueError("cluster_power须在1与2之间")
            self.x = 1.0-(1.0-s)**self.cluster_power
        else:
            raise ValueError("未知grid_kind")
        self.dx = np.diff(self.x)
        self.faces = np.r_[0.0, (self.x[:-1] + self.x[1:]) / 2.0, 1.0]
        self.weights = np.diff(self.faces**2) / 2.0

    def divergence(self, values, face_coefficients, radius, surface_flux):
        """共享面通量差分。

        face_coefficients 为 N 个内部面的系数，surface_flux 为
        归一化坐标 x=1 处的 x a u_x 通量。质量边界取
        -R*hm*(Cs-Ca)，温度边界取 R*(h*(Ta-Ts)-Lv*jw)。
        """
        values = np.asarray(values)
        coefficients = np.asarray(face_coefficients)
        if values.shape != (self.N + 1,) or coefficients.shape != (self.N,):
            raise ValueError("节点值或内部面系数尺寸不匹配")
        if not np.isfinite(radius) or radius <= 0:
            raise ValueError("半径必须为正有限数")
        flux = np.empty(self.N + 2)
        flux[0] = 0.0
        flux[1:-1] = self.faces[1:-1] * coefficients * np.diff(values) / self.dx
        flux[-1] = surface_flux
        return np.diff(flux) / (radius**2 * self.weights)


def material_properties(question: int, T, C):
    """返回 rho, cp, k, D；指数内温度统一转为 K，严禁静默截断 C。"""
    T, C = np.broadcast_arrays(np.asarray(T, dtype=float), np.asarray(C, dtype=float))
    if question not in (1, 2, 3, 4):
        raise ValueError("question 必须为 1、2、3 或 4")
    if np.any(~np.isfinite(T)) or np.any(~np.isfinite(C)):
        raise ValueError("温度或含水率出现非有限数")
    if np.any(C <= 0) or np.any(T + 273.15 <= 0):
        raise ValueError("状态超出正含水率及正绝对温度的模型定义域")
    if question == 1:
        rho = np.full_like(C, 820.0)
        cp = np.full_like(C, 2600.0)
        k = np.full_like(C, 0.36)
        D = 7e-9 * np.exp(-0.89 / C)
    elif question in (2, 3):
        rho = 650.0 + 128.0 * C
        cp = 1450.0 + 2736.0 * C / (C + 1.0)
        k = 0.21 + 0.38 * C / (C + 1.0)
        D = 2.4e-3 * np.exp(-0.45 / C - 3850.0 / (T + 273.15))
    else:
        rho = 760.0 + 90.0 * C
        cp = 1850.0 + 2150.0 * C / (C + 1.0)
        k = 0.12 + 0.20 * C / (C + 1.0)
        D = 4.2e-4 * np.exp(-0.30 / C - 3850.0 / (T + 273.15))
    return rho, cp, k, D


class RadialModel:
    """半离散有效模型；y = [T_0,...,T_N,C_0,...,C_N]。"""

    def __init__(self, question, inputs, N=160, latent=False, fixed_radius=False,
                 boundary_extension="last_hour_mean", latent_heat=2.38e6,
                 ambient_moisture_offset=0.0, gauss_order=3,
                 grid_kind="uniform", cluster_power=1.5, moving_radius=None,
                 D_scale=1.0, hm_scale=1.0, ambient_temperature_offset=0.0,
                 equilibrium_scale=1.0):
        if question not in (1, 2, 3, 4):
            raise ValueError("question 必须为 1、2、3 或 4")
        self.question, self.inputs, self.grid = question, inputs, FVMGrid(N, grid_kind, cluster_power)
        self.latent, self.fixed_radius = bool(latent), bool(fixed_radius)
        self.use_measured_radius = (question == 4 and not fixed_radius) if moving_radius is None else bool(moving_radius)
        if fixed_radius and self.use_measured_radius:
            raise ValueError("固定半径与强制收缩不能同时启用")
        if min(D_scale, hm_scale) <= 0 or equilibrium_scale < 0:
            raise ValueError("扩散/传质倍率须为正，平衡含水率倍率非负")
        self.D_scale = float(D_scale)
        self.hm_scale = float(hm_scale)
        self.ambient_temperature_offset = float(ambient_temperature_offset)
        self.equilibrium_scale = float(equilibrium_scale)
        self.boundary_extension = boundary_extension
        self.latent_heat = float(latent_heat)
        self.ambient_moisture_offset = float(ambient_moisture_offset)
        if not np.isfinite(self.latent_heat) or self.latent_heat <= 0:
            raise ValueError("汽化潜热必须为正有限数")
        if not np.isfinite(self.ambient_moisture_offset):
            raise ValueError("环境浓度偏移必须为有限数")
        self.m = N + 1
        self.initial_state = np.r_[np.full(self.m, T0), np.full(self.m, C0)]
        initial_rho = float(material_properties(question, T0, C0)[0])
        self.initial_dry_density = initial_rho / (1.0 + C0)
        self.rho_d0 = self.initial_dry_density
        self.hm = HM * self.hm_scale
        self.dry_mass = self.initial_dry_density * np.pi * R0**2 * LENGTH
        offsets = [-1, 0, 1]
        block = diags([np.ones(self.m - 1), np.ones(self.m), np.ones(self.m - 1)],
                      offsets, shape=(self.m, self.m), format="csc")
        self.jac_sparsity = bmat([[block, block], [block, block]], format="csc")
        if gauss_order not in (3, 5, 9):
            raise ValueError("gauss_order须为3、5或9")
        nodes, weights = np.polynomial.legendre.leggauss(gauss_order)
        self._gauss_s = ((nodes+1.0)/2.0)[:, None]
        self._gauss_w = (weights/2.0)[:, None]
        self.gauss_order = gauss_order

    def split(self, y):
        y = np.asarray(y)
        if y.shape[0] != 2 * self.m:
            raise ValueError("状态向量维数错误")
        return y[:self.m], y[self.m:]

    def radius(self, t):
        if self.use_measured_radius:
            return self.inputs.radius(t)
        values = np.asarray(t, dtype=float)
        result = np.full_like(values, R0)
        return float(result) if values.ndim == 0 else result

    def dry_density(self, t):
        """由初始标定和几何守恒得出，不再使用局部 rho_eff(C)/(1+C)。"""
        return self.initial_dry_density * (R0 / self.radius(t))**2

    def boundary(self, t, phase=None):
        ta, ca = self.inputs.boundary(t, self.boundary_extension, phase)
        ca = (ca + self.ambient_moisture_offset) * self.equilibrium_scale
        ta = ta + self.ambient_temperature_offset
        if np.any(np.asarray(ca) < 0):
            raise ValueError("偏移后环境有效浓度为负，当前情景不可接受")
        return ta, ca

    def surface_water_flux(self, t, y, phase=None):
        _, C = self.split(y)
        _, ca = self.boundary(t, phase)
        return self.dry_density(t) * self.hm * (C[-1] - ca)

    def rhs(self, t, y, phase=None):
        T, C = self.split(y)
        radius = self.radius(t)
        ta, ca = self.boundary(t, phase)
        rho, cp, k, _ = material_properties(self.question, T, C)
        k_faces = 2 * k[:-1] * k[1:] / (k[:-1] + k[1:])
        Ti = T[:-1][None, :] + self._gauss_s * np.diff(T)[None, :]
        Ci = C[:-1][None, :] + self._gauss_s * np.diff(C)[None, :]
        D_faces = self.D_scale * np.sum(self._gauss_w * material_properties(self.question, Ti, Ci)[3], axis=0)
        moisture_surface_flux = -radius * self.hm * (C[-1] - ca)
        dC = self.grid.divergence(C, D_faces, radius, moisture_surface_flux)
        heat_inward = H * (ta - T[-1])
        if self.latent:
            # 与质量守恒同一个水通量。负通量亦保留其凝结释热符号。
            heat_inward -= self.latent_heat * self.surface_water_flux(t, y, phase)
        dT = self.grid.divergence(T, k_faces, radius, radius * heat_inward) / (rho * cp)
        return np.r_[dT, dC]

    def rhs_conservation_residual(self, t, y, phase=None):
        """半离散平均含水率守恒残差（kg/kg/s），应接近机器精度。"""
        _, C = self.split(y)
        _, dC = self.split(self.rhs(t, y, phase))
        _, ca = self.boundary(t, phase)
        return float(2 * np.dot(self.grid.weights, dC)
                     + 2 * self.hm / self.radius(t) * (C[-1] - ca))


@dataclass
class Sample:
    t: Any
    radius: Any
    r: np.ndarray
    T: np.ndarray
    C: np.ndarray


@dataclass
class _Segment:
    start: float
    end: float
    result: Any
    phase: str


class DryingSolution:
    """分段连续数值解。sample 严格拒绝超出已积分时间范围的请求。"""

    def __init__(self, model, segments, t_critical=None, t_operation=None):
        self.model = model
        self.question = model.question
        self.grid = model.grid
        self.x = model.grid.x.copy()
        self.segments = segments
        self.t_critical = t_critical
        self.t_operation = t_operation
        self.t_end = segments[-1].end
        self.stats = {
            "question": self.question, "N": model.grid.N,
            "nfev": int(sum(s.result.nfev for s in segments)),
            "njev": int(sum(s.result.njev for s in segments)),
            "nlu": int(sum(s.result.nlu for s in segments)),
            "accepted_steps": int(sum(len(s.result.t) - 1 for s in segments)),
            "segments": len(segments), "t_end_s": self.t_end,
            "t_critical_s": t_critical, "t_operation_s": t_operation,
        }

    def state(self, t):
        """返回求解状态；时间数组时形状 (2*(N+1), nt)。"""
        times = np.asarray(t, dtype=float)
        if times.ndim > 1 or np.any(~np.isfinite(times)):
            raise ValueError("时间必须为有限标量或一维数组")
        if np.any(times < 0) or np.any(times > self.t_end):
            raise ValueError(f"只积分到 {self.t_end:g} s；不能外推连续解")
        flat = times.reshape(-1)
        values = np.empty((2 * self.model.m, len(flat)))
        assigned = np.zeros(len(flat), dtype=bool)
        for segment in self.segments:
            mask = (flat >= segment.start) & (flat <= segment.end) & ~assigned
            if np.any(mask):
                values[:, mask] = segment.result.sol(flat[mask])
                assigned[mask] = True
        if not np.all(assigned):
            raise RuntimeError("分段数值解有未覆盖的时间间隙")
        return values[:, 0] if times.ndim == 0 else values

    def sample(self, t, physical_r=None) -> Sample:
        """在已积分时刻采样；物理半径 r 超过当前 R(t) 时返回 NaN。

        默认所有网格节点。给定 physical_r 时，空间使用线性内插，
        不把固定位置误作归一化位置，也不向材料外部外推。
        """
        times = np.asarray(t, dtype=float)
        scalar = times.ndim == 0
        flat = times.reshape(-1)
        y = self.state(flat)
        T, C = self.model.split(y)
        T, C = T.T, C.T
        radius = np.asarray(self.model.radius(flat))
        if physical_r is None:
            radii = radius[:, None] * self.x[None, :]
        else:
            requested = np.atleast_1d(np.asarray(physical_r, dtype=float))
            if requested.ndim != 1 or np.any(~np.isfinite(requested)) or np.any(requested < 0):
                raise ValueError("physical_r 须为非负有限半径的一维数组")
            out_T = np.full((len(flat), len(requested)), np.nan)
            out_C = np.full_like(out_T, np.nan)
            for i, R in enumerate(radius):
                inside = requested <= R
                out_T[i, inside] = np.interp(requested[inside] / R, self.x, T[i])
                out_C[i, inside] = np.interp(requested[inside] / R, self.x, C[i])
            T, C = out_T, out_C
            radii = np.broadcast_to(requested, T.shape).copy()
        if scalar:
            return Sample(float(times), float(radius[0]), radii[0], T[0], C[0])
        return Sample(times.copy(), radius, radii, T, C)

    def mean_C(self, t):
        values = self.model.split(self.state(t))[1]
        return 2 * np.dot(self.model.grid.weights, values)

    def rhs_conservation_residual(self, t):
        times = np.atleast_1d(np.asarray(t, dtype=float))
        states = self.state(times)
        result = np.array([self.model.rhs_conservation_residual(float(time), states[:, i])
                           for i, time in enumerate(times)])
        return float(result[0]) if np.asarray(t).ndim == 0 else result

    def water_balance(self, t):
        """独立用表面出水率梯形积分核验质量（供时间积分加密检查）。

        t 须含 0 且严格递增；integrated_outflow_kg 含采样积分误差，
        不是严格求解误差上界。半离散精确恒等式另见
        rhs_conservation_residual；建议逐步减小审计时间间隔。
        """
        times = np.asarray(t, dtype=float)
        if times.ndim != 1 or len(times) < 2 or times[0] != 0 or np.any(np.diff(times) <= 0):
            raise ValueError("质量审计时间须从 0 开始且严格递增，至少两个时点")
        sample = self.sample(times)
        _, ca = self.model.boundary(times)
        flux = self.model.dry_density(times) * self.model.hm * (sample.C[:, -1] - ca)
        mass_rate = 2 * np.pi * sample.radius * LENGTH * flux
        # 对实测/平台边界跳变拆开积分，保留左右侧各自的端点通量。
        cumulative = np.zeros_like(times)
        for i in range(1, len(times)):
            a, b = times[i - 1], times[i]
            tb = self.model.inputs.boundary_end
            if a < tb < b:
                yt = self.state(tb)
                R = self.model.radius(tb)
                left = 2 * np.pi * R * LENGTH * self.model.surface_water_flux(tb, yt, "observed")
                right = 2 * np.pi * R * LENGTH * self.model.surface_water_flux(tb, yt, "plateau")
                area = (mass_rate[i - 1] + left) * (tb - a) / 2
                area += (right + mass_rate[i]) * (b - tb) / 2
            else:
                fa, fb = mass_rate[i - 1], mass_rate[i]
                if a == tb:
                    fa = 2 * np.pi * self.model.radius(a) * LENGTH * self.model.surface_water_flux(a, self.state(a), "plateau")
                area = (fa + fb) * (b - a) / 2
            cumulative[i] = cumulative[i - 1] + area
        water_mass = self.model.dry_mass * self.mean_C(times)
        residual = water_mass - water_mass[0] + cumulative
        return {"t": times, "water_mass_kg": water_mass,
                "outflow_kg_s": mass_rate, "integrated_outflow_kg": cumulative,
                "residual_kg": residual,
                "relative_to_initial_water": residual / water_mass[0]}


def simulate(question: int, inputs: ModelInputs, N: int = 160,
             method: str = "BDF", t_end: float | None = None,
             stop_at_threshold: bool | None = None, threshold: float = THRESHOLD,
             latent: bool = False, fixed_radius: bool = False,
             boundary_extension: str = "last_hour_mean", latent_heat: float = 2.38e6,
             ambient_moisture_offset: float = 0.0,
             rtol: float = 2e-8, atol: float = 2e-10,
             max_step_early: float = 30.0, max_step_late: float = 600.0,
             gauss_order: int = 3, grid_kind: str = "uniform",
             cluster_power: float = 1.5, moving_radius: bool | None = None,
             D_scale: float = 1.0, hm_scale: float = 1.0,
             ambient_temperature_offset: float = 0.0,
             equilibrium_scale: float = 1.0) -> DryingSolution:
    """运行四问有效模型，隐式积分并在边界平台起点重新启动。

    Q1 默认积分至 1800 s；Q2 默认积分至 72 h；Q3/4 默认在全体节点
    max(C)=0.15 时定位临界时刻，再真正续积到其后的下一整分钟。
    若在 t_end 前未达阈值，则 t_critical/t_operation 为 None，保留
    完整已积分轨迹；调用者不可把该时段末尾视为达标。第4问不外推 R。
    """
    if method not in ("BDF", "Radau"):
        raise ValueError("method 须为 BDF 或 Radau")
    if t_end is None:
        t_end = 1800.0 if question == 1 else 72 * 3600.0
    t_end = float(t_end)
    if not np.isfinite(t_end) or t_end <= 0:
        raise ValueError("t_end 必须为正有限数")
    if stop_at_threshold is None:
        stop_at_threshold = question in (3, 4)
    if not 0 < threshold < C0:
        raise ValueError("threshold 必须介于 0 与初始含水率之间")
    if rtol <= 0 or atol <= 0 or max_step_early <= 0 or max_step_late <= 0:
        raise ValueError("容差及最大步长必须为正")
    model = RadialModel(question, inputs, N, latent, fixed_radius,
                        boundary_extension, latent_heat, ambient_moisture_offset,
                        gauss_order, grid_kind, cluster_power, moving_radius,
                        D_scale, hm_scale, ambient_temperature_offset, equilibrium_scale)
    if model.use_measured_radius and t_end > inputs.radius_end:
        raise ValueError("第4问 t_end 超过实测半径区间；禁止半径外推")

    def event(t, y):
        return float(np.max(model.split(y)[1]) - threshold)

    event.direction = -1
    event.terminal = True
    segments = []
    t_critical = None
    t_operation = None

    def integrate(start, end, y0, phase, detect):
        max_step = max_step_early if phase == "observed" else max_step_late
        result = solve_ivp(lambda t, y: model.rhs(t, y, phase), (start, end), y0,
                           method=method, rtol=rtol, atol=atol, dense_output=True,
                           # 极细网格下默认显式试探会越出物性定义域；仅限制初始试探。
                           first_step=min(1e-6,end-start) if start==0 and N>=2048 else None,
                           jac_sparsity=model.jac_sparsity, max_step=max_step,
                           events=event if detect else None)
        if not result.success:
            raise RuntimeError(f"问题{question}在 {result.t[-1]:.6g} s 积分失败：{result.message}")
        segment = _Segment(float(start), float(result.t[-1]), result, phase)
        segments.append(segment)
        return result

    def advance(start, end, y0, detect):
        """任意续积区间均在边界末时刻拆分，绝不越界评估dense输出。"""
        current, state = float(start), y0
        while current < end:
            before = current < inputs.boundary_end
            bound = min(end, inputs.boundary_end) if before else end
            phase = "observed" if before else "plateau"
            result = integrate(current, bound, state, phase, detect)
            current, state = float(result.t[-1]), result.y[:, -1]
            if detect and len(result.t_events[0]):
                return float(result.t_events[0][0]), state
        return None, state

    hit, state = advance(0.0, t_end, model.initial_state, bool(stop_at_threshold))
    if hit is not None:
        t_critical = hit
        # floor+1确保即使临界时刻恰为整分钟，也真正进入下一整分钟。
        t_operation = float((math.floor(hit / 60.0) + 1) * 60)
        if model.use_measured_radius and t_operation > inputs.radius_end:
            raise RuntimeError("临界时刻位于半径观测末端，下一整分钟无半径数据；不能确认操作时刻")
        advance(hit, t_operation, state, False)
        if np.max(model.split(segments[-1].result.y[:, -1])[1]) >= threshold:
            raise RuntimeError("续积到下一整分钟后未严格达标；请检查非单调环境或延长操作情景")
    solution = DryingSolution(model, segments, t_critical, t_operation)
    solution.stats.update({"method": method, "rtol": rtol, "atol": atol,
                           "latent": latent, "fixed_radius": fixed_radius,
                           "boundary_extension": boundary_extension,
                           "max_step_early_s": max_step_early,
                           "max_step_late_s": max_step_late, "gauss_order": gauss_order,
                           "grid_kind": grid_kind, "cluster_power": cluster_power,
                           "moving_radius": model.use_measured_radius,
                           "D_scale": D_scale, "hm_scale": hm_scale,
                           "ambient_temperature_offset": ambient_temperature_offset,
                           "equilibrium_scale": equilibrium_scale})
    return solution
